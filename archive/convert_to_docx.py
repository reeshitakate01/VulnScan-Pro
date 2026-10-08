import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# Page setup
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

# Styles
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 4):
    heading_style = doc.styles[f'Heading {level}']
    heading_style.font.name = 'Calibri'
    heading_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

# Read the markdown
md_path = os.path.join(os.path.dirname(__file__), 'docs', 'scanner_modules_guide.md')
with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h.strip().strip('`')
        for p in cell.paragraphs:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(10)
    
    # Data rows
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = val.strip().strip('`')
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    
    doc.add_paragraph()

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x2E, 0x3A, 0x4F)

i = 0
in_code = False
code_buffer = []
in_table = False
table_rows = []
table_headers = []

while i < len(lines):
    line = lines[i].rstrip('\r\n')
    
    # Code block detection
    if line.strip().startswith('```'):
        if in_code:
            add_code_block(doc, '\n'.join(code_buffer))
            code_buffer = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    
    if in_code:
        code_buffer.append(line)
        i += 1
        continue
    
    # Table detection
    if '|' in line and line.strip().startswith('|'):
        cols = [c.strip() for c in line.split('|')[1:-1]]
        
        # Check if next line is separator
        if i + 1 < len(lines) and re.match(r'^\s*\|[\s\-|]+\|\s*$', lines[i+1]):
            # This is a header row
            table_headers = cols
            table_rows = []
            i += 2  # skip header and separator
            # Collect data rows
            while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
                row_cols = [c.strip() for c in lines[i].split('|')[1:-1]]
                table_rows.append(row_cols)
                i += 1
            add_table(doc, table_headers, table_rows)
            continue
        else:
            # Data row of existing table or standalone
            i += 1
            continue
    
    # Horizontal rule
    if line.strip() == '---':
        doc.add_paragraph()  # just add spacing
        i += 1
        continue
    
    # Empty line
    if not line.strip():
        i += 1
        continue
    
    # Headings
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=0)
        i += 1
        continue
    if line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=1)
        i += 1
        continue
    if line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=2)
        i += 1
        continue
    
    # Bullet points
    if line.strip().startswith('- '):
        text = line.strip()[2:]
        # Clean markdown formatting
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        p = doc.add_paragraph(style='List Bullet')
        
        # Handle bold parts
        parts = re.split(r'(\*\*.*?\*\*)', line.strip()[2:])
        p.clear()
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                clean = re.sub(r'`(.*?)`', r'\1', part)
                p.add_run(clean)
        
        i += 1
        continue
    
    # Numbered list
    if re.match(r'^\d+\.\s', line.strip()):
        text = re.sub(r'^\d+\.\s', '', line.strip())
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'`(.*?)`', r'\1', text)
        p = doc.add_paragraph(style='List Number')
        
        parts = re.split(r'(\*\*.*?\*\*)', re.sub(r'^\d+\.\s', '', line.strip()))
        p.clear()
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                clean = re.sub(r'`(.*?)`', r'\1', part)
                p.add_run(clean)
        
        i += 1
        continue
    
    # Bold-label lines (like **File:** or **Key Technical Features:**)
    if line.strip().startswith('**'):
        text = line.strip()
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                clean = re.sub(r'`(.*?)`', r'\1', part)
                p.add_run(clean)
        i += 1
        continue
    
    # Regular paragraph
    text = line.strip()
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    
    # Check if next lines continue this paragraph
    full_text = text
    while i + 1 < len(lines):
        next_line = lines[i + 1].rstrip('\r\n')
        if (not next_line.strip() or next_line.startswith('#') or 
            next_line.startswith('-') or next_line.startswith('|') or
            next_line.startswith('```') or next_line.strip().startswith('**') or
            next_line.strip() == '---' or re.match(r'^\d+\.\s', next_line.strip())):
            break
        cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', next_line.strip())
        cleaned = re.sub(r'`(.*?)`', r'\1', cleaned)
        full_text += ' ' + cleaned
        i += 1
    
    doc.add_paragraph(full_text)
    i += 1

# Save
output_path = os.path.join(os.path.dirname(__file__), 'docs', 'Scanner_Modules_Guide.docx')
doc.save(output_path)
print(f"Document saved to: {output_path}")
